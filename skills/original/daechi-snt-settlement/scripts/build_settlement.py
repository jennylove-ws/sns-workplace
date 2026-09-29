"""
대치SNT학원 정산 파일 생성 헬퍼 템플릿.

이 파일은 실행 가능한 완성 스크립트가 아니라, 매번 선생님별로 값을 채워 재사용하는
"템플릿"입니다. 아래 순서로 사용하세요:

1. 이 파일을 복사해 `make_<선생님>.py` 로 이름을 바꾼다.
2. HEADER 영역(폰트/색상/제목/시트 이름/MAX_COL)을 이 선생님에 맞게 수정한다.
3. 반별로 build_table(...)을 호출해 표를 쌓는다 (student 튜플: (이름, 학교, {날짜키:1|"pending"}, [메모]) ).
4. references/teacher_policy.md 에서 방식(기본/60%/중고등구분/합계만)을 확인하고
   그에 맞는 정산내역 수식을 마지막에 작성한다 (policy_*.md 참고).
5. 저장 후 `python3 /mnt/skills/public/xlsx/scripts/recalc.py <파일>.xlsx 90` 으로 재계산·검증한다.

이 파일 자체를 그대로 실행하면 안 됩니다 (예시 변수가 비어 있음).
"""

import copy
import openpyxl
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.comments import Comment
from openpyxl.utils import get_column_letter

# ============== 0. 여기를 선생님/파일에 맞게 수정 ==============
OUT = "/home/claude/build/대치SNT_단가_비율제_정산_XXX.xlsx"
SHEET_TITLE = "XXXT"
TITLE_TEXT = "2026.07 대치SNT학원 XXX 선생님 수업내역"

wb = openpyxl.Workbook()
ws = wb.active
ws.title = SHEET_TITLE

FONT_NAME = "나눔바른고딕 Light"
thin = Side(style="thin", color="000000")
border = Border(left=thin, right=thin, top=thin, bottom=thin)
header_fill = PatternFill("solid", fgColor="D9D9D9")
total_fill = PatternFill("solid", fgColor="F2F2F2")
pending_fill = PatternFill("solid", fgColor="FFFF00")   # 노란색: 미확인(출결 '미' / 학교 미확인)
excluded_fill = PatternFill("solid", fgColor="DDEBF7")  # 연한 블루: 회차 미포함
input_font = Font(name=FONT_NAME, size=11, color="0000FF")  # 하드코딩 입력값(단가 등)
formula_font = Font(name=FONT_NAME, size=11, color="000000")
header_font = Font(name=FONT_NAME, size=11, bold=True)
title_font = Font(name=FONT_NAME, size=16, bold=True)
subtitle_font = Font(name=FONT_NAME, size=12, bold=True)
bold_font = Font(name=FONT_NAME, size=11, bold=True)
center = Alignment(horizontal="center", vertical="center", wrap_text=True)
left = Alignment(horizontal="left", vertical="center")
right = Alignment(horizontal="right", vertical="center")

MONEY_FMT = "#,##0"
KRW_FMT = "\u20a9#,##0_);[RED](\u20a9#,##0)"

# 표에 필요한 최대 열 수: 3(No/이름/학교) + 날짜수 + 5(출석횟수/단가/실수강료/납부방법/실입금액)
# 여러 반이 있으면 그중 가장 date-heavy한 반 기준으로 잡는다.
MAX_COL = 12

ws.merge_cells(f"A1:{get_column_letter(MAX_COL)}1")
title_cell = ws["A1"]
title_cell.value = TITLE_TEXT
title_cell.font = title_font
title_cell.alignment = center
ws.row_dimensions[1].height = 30
ws.freeze_panes = "D1"  # No./이름/학교 열 고정


def style_cell(cell, font=formula_font, align=center, fmt=None, fill=None):
    cell.font = font
    cell.alignment = align
    cell.border = border
    if fmt:
        cell.number_format = fmt
    if fill:
        cell.fill = fill


def build_table(start_row, subtitle, date_labels, students, unit_price):
    """
    date_labels: [("7/7(화)", "7/7"), ...]  (표시용 텍스트, 내부 날짜 키)
    students: [(이름, 학교_or_None, {날짜키: 1 또는 "pending"}), ...]
              4번째 요소로 (이름, 학교, att, "M/D(요일)\\t수강시작") 처럼 메모를 줄 수 있음.
              att 딕셔너리에 값이 없거나 없는 키는 "회차 미포함"(연한 블루)으로 처리됨.
              값이 "pending"이면 "미확인"(노란색)으로 처리됨.
    반환값: (합계행 번호, 실입금액 열 index, 데이터시작행, 데이터끝행)
    """
    n_date = len(date_labels)
    headers = ["No.", "이름", "학교"] + [d[0] for d in date_labels] + \
              ["출석횟수", "단가", "실수강료", "납부방법", "실입금액(2.2%-수수료)"]
    ncols = len(headers)
    dep_col = ncols  # 실입금액은 항상 마지막 열
    dep_letter = get_column_letter(dep_col)
    header_row = start_row + 1
    data_start = header_row + 1
    data_end = data_start + len(students) - 1
    total_row = data_end + 1

    # 표 제목 줄: 학생 수·실입금액 합계를 수식으로 이어붙임 (스크롤 없이 바로 보이게)
    ws.merge_cells(start_row=start_row, start_column=1, end_row=start_row, end_column=ncols)
    st = ws.cell(row=start_row, column=1,
                 value=f'="{subtitle}  (총 {len(students)}명 / 실입금액 합계 "&TEXT({dep_letter}{total_row},"#,##0")&"원)"')
    st.font = subtitle_font
    st.alignment = left
    ws.row_dimensions[start_row].height = 22

    for i, h in enumerate(headers, start=1):
        c = ws.cell(row=header_row, column=i, value=h)
        style_cell(c, font=header_font, fill=header_fill)
    ws.row_dimensions[header_row].height = 30

    row = data_start
    for entry in students:
        name, school, att = entry[0], entry[1], entry[2]
        comment_text = entry[3] if len(entry) > 3 else None

        no_cell = ws.cell(row=row, column=1, value=f"=ROW()-{data_start - 1}")
        name_cell = ws.cell(row=row, column=2, value=name)
        if comment_text:
            name_cell.comment = Comment(comment_text, "koyon")
        school_cell = ws.cell(row=row, column=3, value=school)
        if not school:
            school_cell.fill = pending_fill  # 학교 미확인 -> 노란색

        col = 4
        for _, key in date_labels:
            cell = ws.cell(row=row, column=col)
            val = att.get(key)
            if val == 1:
                cell.value = 1
            elif val == "pending":
                cell.fill = pending_fill      # 출결 '미' -> 노란색
            else:
                cell.fill = excluded_fill     # 회차 미포함 -> 연한 블루
            col += 1

        att_col = col
        ws.cell(row=row, column=att_col,
                value=f"=COUNTA({get_column_letter(4)}{row}:{get_column_letter(3 + n_date)}{row})")
        price_col = att_col + 1
        ws.cell(row=row, column=price_col, value=unit_price)
        fee_col = price_col + 1
        ws.cell(row=row, column=fee_col,
                value=f"={get_column_letter(att_col)}{row}*{get_column_letter(price_col)}{row}")
        pay_col = fee_col + 1
        ws.cell(row=row, column=pay_col, value="카드")
        dep_col_row = pay_col + 1
        ws.cell(row=row, column=dep_col_row,
                value=f'=IF({get_column_letter(pay_col)}{row}="카드",0.978*{get_column_letter(fee_col)}{row},{get_column_letter(fee_col)}{row})')

        for cidx in range(1, ncols + 1):
            style_cell(ws.cell(row=row, column=cidx))
        ws.cell(row=row, column=2).alignment = left
        ws.cell(row=row, column=3).alignment = left
        ws.cell(row=row, column=price_col).font = input_font  # 단가 = 하드코딩 입력값 -> 파란 글씨
        ws.cell(row=row, column=fee_col).number_format = MONEY_FMT
        ws.cell(row=row, column=dep_col_row).number_format = KRW_FMT
        row += 1

    # 합계 행
    ws.merge_cells(start_row=total_row, start_column=1, end_row=total_row, end_column=dep_col - 1)
    lbl = ws.cell(row=total_row, column=1, value="합   계")
    style_cell(lbl, font=bold_font, fill=total_fill)
    tot = ws.cell(row=total_row, column=dep_col,
                   value=f"=SUM({dep_letter}{data_start}:{dep_letter}{data_end})")
    style_cell(tot, font=bold_font, fmt=KRW_FMT, fill=total_fill, align=right)
    ws.row_dimensions[total_row].height = 24

    widths = [6, 10, 15] + [11] * n_date + [10, 10, 12, 10, 18]
    for i, w in enumerate(widths, start=1):
        letter = get_column_letter(i)
        cur = ws.column_dimensions[letter].width
        if cur is None or w > cur:
            ws.column_dimensions[letter].width = w

    return total_row, dep_col, data_start, data_end


# ============== 1. 반별 표 생성 (예시 — 실제 값으로 교체) ==============
# dates_1 = [("7/7(화)", "7/7"), ("7/9(목)", "7/9")]
# students_1 = [
#     ("홍길동", "대치고", {"7/7": 1, "7/9": 1}),
#     ("김철수(신)", "역삼고", {"7/9": 1}, "7/9(목)\t수강시작"),
# ]
# total_row_1, dep_col_1, ds_1, de_1 = build_table(3, "① [반이름]", dates_1, students_1, 62500)

# ============== 2. 정산내역 ==============
# references/teacher_policy.md 에서 방식 확인 후 policy_*.md 의 공식을 그대로 이식.
# 아래는 "기본(50%)" 예시:
#
# dep_letter_1 = get_column_letter(dep_col_1)
# SETT_TITLE_ROW = total_row_1 + 2
# LABEL_END, VAL_START, VAL_END = 6, 7, 11  # 정산내역은 항상 표 너비와 무관하게 좁게(A~K열)
# ... (내역/금액 헤더, ①~④, 지급액 행 작성 — 이전 대화의 make_*.py 참고)

# ============== 3. 저장 ==============
# wb.save(OUT)
# print("saved", OUT)
