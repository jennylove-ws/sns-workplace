"""
classes.json -> HTML용 시간표 엑셀 (반명 / 요일-시간 / 강의실), 정규+클리닉 각각 별도 행.
같은 강의실·같은 요일에서 시간이 겹치는 클리닉끼리는 자동으로 한 행에 합쳐서
반명을 " / "로 나열한다 (merge_utils.merge_sessions 참고).
검증된 파이썬 소스에서 직접 생성.
"""
import json, os, sys
sys.path.insert(0, os.path.dirname(__file__))
import openpyxl
from openpyxl.styles import Font, Alignment
from merge_utils import merge_sessions

HERE = os.path.dirname(__file__)
DATA = json.load(open(os.path.join(HERE, "..", "data", "classes.json"), encoding="utf-8"))
TEACHERS = DATA["teachers"]
CLASSES = DATA["classes"]

DAY_ORDER = {"월": 0, "화": 1, "수": 2, "목": 3, "금": 4, "토": 5, "일": 6}

merged = merge_sessions(CLASSES, TEACHERS)

rows = []  # (반명, 요일/시간, 강의실, sort_key)
for m in merged:
    name = " / ".join(m["names"])
    day = m["day"]
    if m["is_merged"]:
        prefix = f"{m['kind']} 합반" if m["kind"] != "정규" else "합반"
        label = f"({day}) {prefix} {m['start']}-{m['end']}"
    elif m["kind"] != "정규":
        label = f"({day}) {m['kind']} {m['start']}-{m['end']}"
    else:
        label = f"{day} {m['start']}-{m['end']}"
    room_disp = f"A관 {m['room']}"
    sort_key = (DAY_ORDER.get(day, 9), m["startMin"], 1 if m["is_clinic"] else 0)
    rows.append((name, label, room_disp, sort_key))

# 강의실 미정인 세션도 표에 넣는다 (merge_sessions는 미정을 건너뛰므로 별도로 채워줌)
for c in CLASSES:
    teacher = TEACHERS.get(c["teacher_id"], c["teacher_id"])
    full_name = f"{c['name']} ({teacher}T)"
    for s in c["sessions"]:
        if s["room"] != "미정":
            continue
        is_clinic = bool(s.get("clinic") or s.get("clinic_union"))
        day = s["day"]
        label = f"({day}) 클리닉 {s['start']}-{s['end']}" if is_clinic else f"{day} {s['start']}-{s['end']}"
        h, mnt = s["start"].split(":")
        sort_key = (DAY_ORDER.get(day, 9), int(h) * 60 + int(mnt), 1 if is_clinic else 0)
        rows.append((full_name, label, "미정", sort_key))

rows.sort(key=lambda r: r[3])

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "HTML용_시간표"

header_font = Font(bold=True)
ws.append(["반명", "요일/시간", "강의실"])
for c in range(1, 4):
    ws.cell(row=1, column=c).font = header_font

for name, label, room, _ in rows:
    ws.append([name, label, room])

ws.column_dimensions["A"].width = 50
ws.column_dimensions["B"].width = 30
ws.column_dimensions["C"].width = 14
for row in ws.iter_rows(min_row=2):
    for cell in row:
        cell.alignment = Alignment(vertical="center", wrap_text=True)

out_path = os.path.join(HERE, "..", "output", "HTML용_시간표.xlsx")
wb.save(out_path)
print("done:", out_path, "rows:", len(rows))
