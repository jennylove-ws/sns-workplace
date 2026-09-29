"""
클래스DB(사실상 classes.json)에서 강의실 배정 문서(docx)를 생성한다.
소스: ../data/classes.json
출력: ../output/강의실_배정.docx
"""
import json, sys, os
sys.path.insert(0, os.path.dirname(__file__))
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

DATA = json.load(open(os.path.join(os.path.dirname(__file__), "..", "data", "classes.json"), encoding="utf-8"))
TEACHERS = DATA["teachers"]
CLASSES = DATA["classes"]

SUBJECT_ORDER = ["수학", "국어", "영어", "약술논술", "과학", "수원외고", "외대부고"]

def fmt_session(s):
    day = f"({s['day']})" if s.get("alt") or s.get("clinic_union") else s["day"]
    tag = ""
    if s.get("clinic"):
        tag = " 클리닉"
    if s.get("clinic_union"):
        tag = " 클리닉연합"
    if s.get("note"):
        tag += f" ({s['note']})"
    return f"{day} {s['start']}-{s['end']}{tag}"

doc = Document()
title = doc.add_heading("강의실 배정 안내 (자동 생성)", level=1)
p = doc.add_paragraph("※ 이 문서는 classes.json(클래스DB)에서 자동 생성되었습니다. 반이 바뀌면 classes.json만 수정하고 다시 실행하세요.")
p.runs[0].italic = True
p.runs[0].font.size = Pt(9)

for subject in SUBJECT_ORDER:
    subject_classes = [c for c in CLASSES if c["subject"] == subject]
    if not subject_classes:
        continue
    doc.add_heading(f"[ {subject} ]", level=2)
    for c in subject_classes:
        teacher = TEACHERS.get(c["teacher_id"], c["teacher_id"])
        head = doc.add_paragraph()
        run = head.add_run(f"■ {c['name']} ({teacher}T)")
        run.bold = True

        pricing = c["pricing"]
        if pricing["type"] == "per_session":
            if pricing.get('unit_fee') is None:
                price_txt = "수강료 정보 확인 필요"
            else:
                price_txt = f"회당 {pricing['unit_fee']:,}원 × {pricing['count']}강 = {pricing['total']:,}원"
        elif pricing['type'] == 'flat_monthly':
            price_txt = f"1개월 정액 {pricing['monthly_fee']:,}원 ({pricing['count']}강)"
        elif pricing['type'] == 'none':
            price_txt = pricing.get('note', '강의실만 대여 - 수강료 없음')
        else:
            price_txt = "수강료 정보 확인 필요"
        doc.add_paragraph(f"-수강료 : {price_txt}")

        # group sessions by room
        rooms = {}
        for s in c["sessions"]:
            rooms.setdefault(s["room"], []).append(s)
        for room, sess_list in rooms.items():
            times = ", ".join(fmt_session(s) for s in sess_list)
            doc.add_paragraph(f"-{times}  →  A관 {room}")

doc.save(os.path.join(os.path.dirname(__file__), "..", "output", "강의실_배정.docx"))
print("done: 강의실_배정.docx")
